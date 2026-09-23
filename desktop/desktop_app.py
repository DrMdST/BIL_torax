"""
NeuroRad Analytics - Desktop Edition
NRRD Radiomics Feature Extractor

Desktop application for Windows 10 (and cross-platform).
Builds into a standalone .exe via PyInstaller - no installation required.

All radiomic features are computed ONLY from the region of interest (ROI)
defined by the .seg.nrrd segmentation mask. Voxels outside the mask are
never included in any calculation.

Usage:
    python desktop_app.py

Build exe:
    pyinstaller --onefile --windowed --name NeuroRad desktop_app.py
"""

import os
import sys
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from typing import Optional

import numpy as np

try:
    import nrrd
except ImportError:
    print("Installing pynrrd...")
    os.system(f"{sys.executable} -m pip install pynrrd")
    import nrrd

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

try:
    from PIL import Image, ImageTk
except ImportError:
    print("Installing Pillow...")
    os.system(f"{sys.executable} -m pip install Pillow")
    from PIL import Image, ImageTk


# ---------------------------------------------------------------------------
# Image processing
# ---------------------------------------------------------------------------

def to_gray(data: np.ndarray) -> np.ndarray:
    """Convert RGB/RGBA data to grayscale, or return as-is if already grayscale."""
    if data.ndim == 2:
        return data.astype(np.float32)
    if data.ndim == 3 and data.shape[-1] in (3, 4):
        rgb = data[..., :3].astype(np.float32)
        return 0.299 * rgb[..., 0] + 0.587 * rgb[..., 1] + 0.114 * rgb[..., 2]
    if data.ndim == 4 and data.shape[-1] in (3, 4):
        rgb = data[..., :3].astype(np.float32)
        return 0.299 * rgb[..., 0] + 0.587 * rgb[..., 1] + 0.114 * rgb[..., 2]
    return data.astype(np.float32).squeeze()


def normalize_slice(slice_data: np.ndarray) -> np.ndarray:
    """Normalize to 0-255 uint8 for visualization."""
    mn, mx = float(slice_data.min()), float(slice_data.max())
    if mx - mn < 1e-8:
        return np.zeros_like(slice_data, dtype=np.uint8)
    return ((slice_data - mn) / (mx - mn) * 255).astype(np.uint8)


def slice_to_image(slice_data: np.ndarray, mask_slice: Optional[np.ndarray] = None) -> Image.Image:
    """Convert a 2D slice to a PIL Image, overlaying mask if provided."""
    norm = normalize_slice(slice_data)
    if mask_slice is not None:
        mask_bin = (mask_slice > 0).astype(np.uint8)
        red_overlay = np.stack([
            np.where(mask_bin > 0, (norm.astype(np.int32) * 2 // 5 + 200).clip(0, 255).astype(np.uint8), norm),
            np.where(mask_bin > 0, (norm.astype(np.int32) * 2 // 5 + 50).clip(0, 255).astype(np.uint8), norm),
            np.where(mask_bin > 0, (norm.astype(np.int32) * 2 // 5 + 50).clip(0, 255).astype(np.uint8), norm),
        ], axis=-1)
        return Image.fromarray(red_overlay, 'RGB')
    return Image.fromarray(np.stack([norm, norm, norm], axis=-1), 'RGB')


# ---------------------------------------------------------------------------
# Feature extraction — ALL features computed ONLY within the mask ROI
# ---------------------------------------------------------------------------

def compute_first_order(masked_values: np.ndarray) -> dict:
    """First-order statistics from image intensities WITHIN the mask only."""
    if masked_values.size == 0:
        return {k: 0.0 for k in [
            'Mean', 'Std', 'Variance', 'Skewness', 'Kurtosis',
            'Min', 'Max', 'Range', 'Median',
            'P10', 'P25', 'P75', 'P90', 'Entropy'
        ]}

    mu = float(masked_values.mean())
    sigma = float(masked_values.std())

    if sigma > 1e-8:
        skew = float(((masked_values - mu) / sigma) ** 3).mean())
        kurt = float(((masked_values - mu) / sigma) ** 4).mean() - 3.0
    else:
        skew = 0.0
        kurt = 0.0

    hist, _ = np.histogram(masked_values, bins=32)
    hist = hist / hist.sum()
    hist = hist[hist > 0]
    ent = float(-(hist * np.log2(hist)).sum())

    return {
        'Mean': mu, 'Std': sigma, 'Variance': float(masked_values.var()),
        'Skewness': skew, 'Kurtosis': kurt,
        'Min': float(masked_values.min()), 'Max': float(masked_values.max()),
        'Range': float(masked_values.max() - masked_values.min()),
        'Median': float(np.median(masked_values)),
        'P10': float(np.percentile(masked_values, 10)),
        'P25': float(np.percentile(masked_values, 25)),
        'P75': float(np.percentile(masked_values, 75)),
        'P90': float(np.percentile(masked_values, 90)),
        'Entropy': ent,
    }


def compute_glcm_masked(img_slice: np.ndarray, mask_slice: np.ndarray, levels: int = 8) -> dict:
    """
    GLCM texture features computed ONLY from pixel pairs where BOTH pixels
    are inside the mask. Gray-level quantization uses the range of the masked
    region only.
    """
    binary = mask_slice > 0
    if binary.sum() == 0:
        return {k: 0.0 for k in ['GLCM_Contrast', 'GLCM_Correlation', 'GLCM_Energy', 'GLCM_Homogeneity']}

    masked_vals = img_slice[binary]
    mn, mx = float(masked_vals.min()), float(masked_vals.max())
    rng = mx - mn if mx > mn else 1.0
    quantized = np.clip(((img_slice - mn) / rng * (levels - 1)).astype(np.int32), 0, levels - 1)

    glcm = np.zeros((levels, levels), dtype=np.float64)
    h, w = quantized.shape
    count = 0
    for y in range(h):
        for x in range(w - 1):
            if binary[y, x] and binary[y, x + 1]:
                a, b = quantized[y, x], quantized[y, x + 1]
                glcm[a, b] += 1
                glcm[b, a] += 1
                count += 2
    if count > 0:
        glcm /= count
    else:
        return {k: 0.0 for k in ['GLCM_Contrast', 'GLCM_Correlation', 'GLCM_Energy', 'GLCM_Homogeneity']}

    contrast, energy, homogeneity = 0.0, 0.0, 0.0
    mu_i, mu_j = 0.0, 0.0
    for i in range(levels):
        for j in range(levels):
            p = glcm[i, j]
            contrast += p * (i - j) ** 2
            energy += p * p
            homogeneity += p / (1 + (i - j) ** 2)
            mu_i += i * p
            mu_j += j * p

    sig_i, sig_j = 0.0, 0.0
    for i in range(levels):
        for j in range(levels):
            p = glcm[i, j]
            sig_i += (i - mu_i) ** 2 * p
            sig_j += (j - mu_j) ** 2 * p
    sig_i = np.sqrt(sig_i)
    sig_j = np.sqrt(sig_j)

    correlation = 0.0
    if sig_i > 0 and sig_j > 0:
        for i in range(levels):
            for j in range(levels):
                p = glcm[i, j]
                correlation += (i - mu_i) * (j - mu_j) * p / (sig_i * sig_j)

    return {
        'GLCM_Contrast': float(contrast),
        'GLCM_Correlation': float(correlation),
        'GLCM_Energy': float(energy),
        'GLCM_Homogeneity': float(homogeneity),
    }


def compute_morphology(mask_slice: np.ndarray) -> dict:
    """Morphological/shape features from the binary mask geometry."""
    binary = mask_slice > 0
    area = int(binary.sum())
    if area == 0:
        return {'Mask_Area': 0, 'Mask_Perimeter': 0, 'Mask_Eccentricity': 0, 'Mask_Extent': 0}

    rows = np.any(binary, axis=1)
    cols = np.any(binary, axis=0)
    rmin, rmax = np.where(rows)[0][[0, -1]]
    cmin, cmax = np.where(cols)[0][[0, -1]]
    bbox_w = cmax - cmin + 1
    bbox_h = rmax - rmin + 1

    perimeter = 0
    h, w = binary.shape
    for y in range(h):
        for x in range(w):
            if binary[y, x]:
                neighbors = 0
                if y > 0 and binary[y - 1, x]: neighbors += 1
                if y < h - 1 and binary[y + 1, x]: neighbors += 1
                if x > 0 and binary[y, x - 1]: neighbors += 1
                if x < w - 1 and binary[y, x + 1]: neighbors += 1
                if neighbors < 4:
                    perimeter += 1

    major = max(bbox_w, bbox_h)
    minor = min(bbox_w, bbox_h)
    eccentricity = float(np.sqrt(1 - (minor / major) ** 2)) if major > 0 else 0.0
    extent = area / (bbox_w * bbox_h) if bbox_w * bbox_h > 0 else 0.0

    return {
        'Mask_Area': area, 'Mask_Perimeter': perimeter,
        'Mask_Eccentricity': eccentricity, 'Mask_Extent': extent,
    }


def extract_features_slice(img_slice: np.ndarray, mask_slice: np.ndarray) -> dict:
    """
    Extract ALL radiomic features for a single slice.

    The mask defines the region of interest (ROI). Only image voxels where
    the mask is non-zero are used for first-order and GLCM calculations.
    Shape features are derived from the mask geometry itself.
    """
    binary = mask_slice > 0
    masked_values = img_slice[binary]

    first_order = compute_first_order(masked_values)
    glcm = compute_glcm_masked(img_slice, mask_slice)
    morph = compute_morphology(mask_slice)

    return {
        'Voxels_in_mask': int(masked_values.size),
        **first_order,
        **glcm,
        **morph,
    }


# ---------------------------------------------------------------------------
# Main Application
# ---------------------------------------------------------------------------

class NeuroRadApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("NeuroRad Analytics - NRRD Radiomics Feature Extractor")
        self.root.geometry("900x700")
        self.root.configure(bg='#0f172a')

        self.image_path: Optional[str] = None
        self.mask_path: Optional[str] = None
        self.image_data: Optional[np.ndarray] = None
        self.mask_data: Optional[np.ndarray] = None
        self.image_header = None
        self.mask_header = None
        self.results: list[dict] = []
        self.processed_image: Optional[Image.Image] = None
        self.processed_mask: Optional[Image.Image] = None
        self.output_dir: Optional[str] = None

        self._build_ui()

    def _build_ui(self):
        style = ttk.Style()
        style.theme_use('clam')
        style.configure('TFrame', background='#0f172a')
        style.configure('TLabel', background='#0f172a', foreground='#e2e8f0', font=('Segoe UI', 10))
        style.configure('Title.TLabel', background='#0f172a', foreground='#22d3ee', font=('Segoe UI', 16, 'bold'))
        style.configure('Sub.TLabel', background='#0f172a', foreground='#94a3b8', font=('Segoe UI', 10))
        style.configure('TButton', font=('Segoe UI', 10), padding=8)
        style.configure('Accent.TButton', font=('Segoe UI', 11, 'bold'), padding=10)

        # Header
        header = ttk.Frame(self.root)
        header.pack(fill='x', padx=20, pady=(15, 10))
        ttk.Label(header, text="NeuroRad Analytics", style='Title.TLabel').pack(side='left')
        ttk.Label(header, text="  NRRD Radiomics Feature Extractor", style='Sub.TLabel').pack(side='left', pady=(8, 0))

        sep = ttk.Separator(self.root)
        sep.pack(fill='x', padx=20, pady=5)

        # File selection
        files_frame = ttk.Frame(self.root)
        files_frame.pack(fill='x', padx=20, pady=10)

        # Image file
        img_frame = ttk.LabelFrame(files_frame, text="Image File (.nrrd)", padding=15)
        img_frame.pack(side='left', fill='x', expand=True, padx=(0, 10))
        self.img_label = ttk.Label(img_frame, text="No file selected", style='Sub.TLabel')
        self.img_label.pack(anchor='w', pady=(0, 10))
        ttk.Button(img_frame, text="Browse .nrrd", command=self._browse_image).pack(anchor='w')

        # Mask file
        mask_frame = ttk.LabelFrame(files_frame, text="Segmentation Mask (.seg.nrrd)", padding=15)
        mask_frame.pack(side='right', fill='x', expand=True, padx=(10, 0))
        self.mask_label = ttk.Label(mask_frame, text="No file selected", style='Sub.TLabel')
        self.mask_label.pack(anchor='w', pady=(0, 10))
        ttk.Button(mask_frame, text="Browse .seg.nrrd", command=self._browse_mask).pack(anchor='w')

        # Output dir
        out_frame = ttk.LabelFrame(self.root, text="Output Directory", padding=15)
        out_frame.pack(fill='x', padx=20, pady=10)
        self.out_label = ttk.Label(out_frame, text="Default: same as image file", style='Sub.TLabel')
        self.out_label.pack(side='left', pady=(0, 10))
        ttk.Button(out_frame, text="Browse", command=self._browse_output).pack(side='right')

        # Process button
        btn_frame = ttk.Frame(self.root)
        btn_frame.pack(fill='x', padx=20, pady=10)
        self.process_btn = ttk.Button(btn_frame, text="Extract Features", style='Accent.TButton', command=self._process)
        self.process_btn.pack(pady=5)

        # Progress
        self.progress = ttk.Progressbar(self.root, mode='indeterminate')

        # Status
        self.status_label = ttk.Label(self.root, text="", style='Sub.TLabel')
        self.status_label.pack(pady=5)

        # Preview
        self.preview_frame = ttk.LabelFrame(self.root, text="Processed Image Preview", padding=10)
        self.preview_frame.pack(fill='both', expand=True, padx=20, pady=10)
        self.preview_canvas = tk.Canvas(self.preview_frame, bg='#1e293b', highlightthickness=0)
        self.preview_canvas.pack(fill='both', expand=True)

    def _browse_image(self):
        path = filedialog.askopenfilename(
            title="Select Image File",
            filetypes=[("NRRD files", "*.nrrd"), ("All files", "*.*")]
        )
        if path:
            self.image_path = path
            self.img_label.config(text=os.path.basename(path))

    def _browse_mask(self):
        path = filedialog.askopenfilename(
            title="Select Segmentation Mask",
            filetypes=[("NRRD seg files", "*.seg.nrrd"), ("NRRD files", "*.nrrd"), ("All files", "*.*")]
        )
        if path:
            self.mask_path = path
            self.mask_label.config(text=os.path.basename(path))

    def _browse_output(self):
        path = filedialog.askdirectory(title="Select Output Directory")
        if path:
            self.output_dir = path
            self.out_label.config(text=path)

    def _process(self):
        if not self.image_path or not self.mask_path:
            messagebox.showwarning("Missing Files", "Please select both an image file and a segmentation mask.")
            return

        self.process_btn.config(state='disabled')
        self.progress.pack(pady=5, padx=20, fill='x')
        self.progress.start()
        self.status_label.config(text="Loading files...")

        thread = threading.Thread(target=self._run_processing, daemon=True)
        thread.start()

    def _run_processing(self):
        try:
            self.root.after(0, lambda: self.status_label.config(text="Loading NRRD files..."))
            self.image_data, self.image_header = nrrd.read(self.image_path)
            self.mask_data, self.mask_header = nrrd.read(self.mask_path)

            self.root.after(0, lambda: self.status_label.config(text="Converting to grayscale..."))
            img_gray = to_gray(self.image_data)
            mask_gray = to_gray(self.mask_data)

            depth = min(img_gray.shape[0] if img_gray.ndim >= 3 else 1, 50)

            self.root.after(0, lambda: self.status_label.config(text=f"Extracting features from {depth} slices..."))

            all_results = []
            for i in range(depth):
                if img_gray.ndim >= 3:
                    img_slice = img_gray[i]
                else:
                    img_slice = img_gray

                if mask_gray.ndim >= 3:
                    mask_slice = mask_gray[i]
                else:
                    mask_slice = mask_gray

                if img_slice.shape != mask_slice.shape:
                    mask_slice = mask_slice[:img_slice.shape[0], :img_slice.shape[1]]

                features = extract_features_slice(img_slice, mask_slice)
                features['Slice'] = i
                all_results.append(features)

                if i % 5 == 0:
                    pct = int((i / depth) * 100)
                    self.root.after(0, lambda p=pct: self.status_label.config(text=f"Processing: {p}%"))

            self.results = all_results

            # Generate processed images (middle slice)
            mid = depth // 2
            if img_gray.ndim >= 3:
                mid_img = img_gray[mid]
                mid_mask = mask_gray[mid] if mask_gray.ndim >= 3 else mask_gray
            else:
                mid_img = img_gray
                mid_mask = mask_gray

            self.root.after(0, lambda: self.status_label.config(text="Generating processed images..."))
            self.processed_image = slice_to_image(mid_img, mid_mask)
            self.processed_mask = slice_to_image(mid_mask, None)

            # Save outputs
            self.root.after(0, lambda: self.status_label.config(text="Saving outputs..."))
            out_dir = self.output_dir or os.path.dirname(self.image_path)

            xlsx_path = os.path.join(out_dir, "radiomics_results.xlsx")
            self._save_xlsx(xlsx_path)

            img_png_path = os.path.join(out_dir, "processed_image.png")
            self.processed_image.save(img_png_path)

            mask_png_path = os.path.join(out_dir, "processed_mask.png")
            self.processed_mask.save(mask_png_path)

            self.root.after(0, lambda: self._on_done(xlsx_path, img_png_path, mask_png_path))

        except Exception as e:
            self.root.after(0, lambda: self._on_error(str(e)))

    def _save_xlsx(self, path: str):
        wb = Workbook()
        ws = wb.active
        ws.title = "Features"

        if not self.results:
            return

        headers = list(self.results[0].keys())
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        for row_idx, row_data in enumerate(self.results, 2):
            for col_idx, key in enumerate(headers, 1):
                val = row_data.get(key, 0)
                if isinstance(val, float):
                    val = round(val, 6)
                ws.cell(row=row_idx, column=col_idx, value=val)

        for col in range(1, len(headers) + 1):
            max_len = max(len(str(ws.cell(row=r, column=col).value or '')) for r in range(1, len(self.results) + 2))
            ws.column_dimensions[chr(64 + col) if col <= 26 else 'A'].width = min(max_len + 2, 30)

        wb.save(path)

    def _on_done(self, xlsx_path: str, img_path: str, mask_path: str):
        self.progress.stop()
        self.progress.pack_forget()
        self.process_btn.config(state='normal')
        self.status_label.config(text=f"Done! {len(self.results)} slices processed.")

        self._show_preview(self.processed_image)

        messagebox.showinfo(
            "Complete",
            f"Analysis complete!\n\n"
            f"Features extracted from {len(self.results)} slices.\n"
            f"All features computed within the mask region of interest.\n\n"
            f"Output files:\n"
            f"  Excel: {xlsx_path}\n"
            f"  Image: {img_path}\n"
            f"  Mask:  {mask_path}"
        )

    def _on_error(self, err: str):
        self.progress.stop()
        self.progress.pack_forget()
        self.process_btn.config(state='normal')
        self.status_label.config(text=f"Error: {err}")
        messagebox.showerror("Error", f"Processing failed:\n{err}")

    def _show_preview(self, pil_img: Image.Image):
        for w in self.preview_canvas.winfo_children():
            w.destroy()

        canvas_w = max(self.preview_canvas.winfo_width(), 400)
        canvas_h = max(self.preview_canvas.winfo_height(), 300)
        img_w, img_h = pil_img.size
        scale = min(canvas_w / img_w, canvas_h / img_h, 1.0)
        new_size = (int(img_w * scale), int(img_h * scale))
        resized = pil_img.resize(new_size, Image.LANCZOS)
        photo = ImageTk.PhotoImage(resized)

        label = tk.Label(self.preview_canvas, image=photo, bg='#1e293b')
        label.image = photo
        label.pack(expand=True)


def main():
    root = tk.Tk()
    app = NeuroRadApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
